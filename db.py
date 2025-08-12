import os
from typing import List, Optional, Tuple
from uuid import uuid4

# Tenta acessar st.secrets no Streamlit
try:
    import streamlit as st
except Exception:
    st = None

import socket
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from dotenv import load_dotenv
load_dotenv()

# --- Garante psycopg instalado antes de importar submódulos ---
try:
    import psycopg  # v3
except Exception:
    # fallback (temporário) para o Cloud, caso ignore o requirements
    import sys, subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "psycopg-binary==3.2.1"])
    import psycopg

# tuple_row é opcional dependendo do wheel
try:
    from psycopg.rows import tuple_row
except Exception:
    tuple_row = None

# ------------------ Sanitização e obtenção da URL ------------------
def _sanitize_neon_url(raw: str) -> str:
    """
    - Remove aspas/espacos
    - Remove channel_binding
    - Remove ':PORT' inválido no netloc
    """
    url = (raw or "").strip().strip('"').strip("'")
    if not url:
        return ""

    parts = urlsplit(url)
    netloc = parts.netloc

    # separa userinfo e host:port
    if "@" in netloc:
        userinfo, hostport = netloc.rsplit("@", 1)
    else:
        userinfo, hostport = "", netloc

    # se houver ':algo' e esse 'algo' não for numérico, remove
    if ":" in hostport:
        host, maybe_port = hostport.rsplit(":", 1)
        if not maybe_port.isdigit():
            hostport = host  # descarta porta inválida
    # remonta netloc
    netloc = f"{userinfo + '@' if userinfo else ''}{hostport}"

    # remove 'channel_binding' do query
    q = dict(parse_qsl(parts.query, keep_blank_values=True))
    q.pop("channel_binding", None)

    cleaned = parts._replace(netloc=netloc, query=urlencode(q))
    return urlunsplit(cleaned)

def _get_conn_str() -> str:
    """
    Prioridade:
      1) st.secrets["NEON_DATABASE_URL"] (Streamlit Cloud)
      2) env var NEON_DATABASE_URL (.env/local)
    + Sanitiza a URL e valida DNS do host.
    """
    url = None
    if st is not None:
        try:
            url = st.secrets.get("NEON_DATABASE_URL", None)
        except Exception:
            pass
    if not url:
        url = os.getenv("NEON_DATABASE_URL", "")

    url = _sanitize_neon_url(url)
    if not url:
        raise RuntimeError("NEON_DATABASE_URL não configurada em st.secrets ou .env")

    # Valida DNS do host (sem acessar parts.port pra não estourar erro)
    parts = urlsplit(url)
    host = parts.hostname
    if not host:
        raise RuntimeError(f"NEON_DATABASE_URL inválida: {url!r}")

    try:
        # usa porta padrão 5432 só para teste de resolução
        socket.getaddrinfo(host, 5432)
    except Exception as e:
        raise RuntimeError(
            f"Não consegui resolver o host '{host}'. "
            "Confira a URL do Neon (use o endpoint '-pooler') e remova placeholders/linhas extras.\n"
            f"Erro: {e}"
        )
    return url

PG_URL = _get_conn_str()

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

# ===================== Conexão =====================
def conectar_banco() -> psycopg.Connection:
    """Retorna conexão Postgres (Neon)."""
    return psycopg.connect(PG_URL, row_factory=tuple_row)

# ===================== Inicialização com FTS (Postgres) =====================
def inicializa_banco():
    """
    Cria tabelas e FTS (tsvector + GIN) no Postgres.
    """
    with conectar_banco() as conn, conn.cursor() as c:
        c.execute("""
        CREATE TABLE IF NOT EXISTS biblioteca (
            id UUID PRIMARY KEY,
            titulo TEXT NOT NULL,
            descricao TEXT,
            tag TEXT,
            arquivo TEXT NOT NULL,
            data_envio TIMESTAMPTZ NOT NULL DEFAULT NOW()
        );
        """)
        c.execute("""
        CREATE TABLE IF NOT EXISTS biblioteca_text (
            id UUID PRIMARY KEY REFERENCES biblioteca(id) ON DELETE CASCADE,
            conteudo TEXT NOT NULL
        );
        """)
        c.execute("""
        CREATE TABLE IF NOT EXISTS conversas (
            id BIGSERIAL PRIMARY KEY,
            usuario TEXT,
            pergunta TEXT,
            resposta TEXT,
            data TIMESTAMPTZ NOT NULL DEFAULT NOW()
        );
        """)
        c.execute("""
        CREATE TABLE IF NOT EXISTS usuarios (
            id UUID PRIMARY KEY,
            nome_completo TEXT NOT NULL,
            cargo TEXT,
            setor TEXT,
            matricula TEXT,
            email TEXT UNIQUE NOT NULL,
            telefone TEXT,
            data_admissao TEXT,
            tipo_contrato TEXT,
            unidade TEXT,
            senha_hash BYTEA NOT NULL,
            perfil TEXT DEFAULT 'usuario',
            status TEXT DEFAULT 'pendente',
            ultimo_acesso TEXT
        );
        """)
        # ---------- FTS ----------
        c.execute("""
        CREATE TABLE IF NOT EXISTS biblioteca_fts (
            id UUID PRIMARY KEY REFERENCES biblioteca(id) ON DELETE CASCADE,
            doc tsvector
        );
        """)
        c.execute("CREATE INDEX IF NOT EXISTS idx_biblioteca_fts_doc ON biblioteca_fts USING GIN (doc);")
        # ---------- FTS: tabela e índice ----------
        c.execute("""
        CREATE TABLE IF NOT EXISTS biblioteca_fts (
            id UUID PRIMARY KEY REFERENCES biblioteca(id) ON DELETE CASCADE,
            doc tsvector
        );
        """)
        c.execute("CREATE INDEX IF NOT EXISTS idx_biblioteca_fts_doc ON biblioteca_fts USING GIN (doc);")

        # ---------- PATCH: remove funções/triggers antigos ----------
        c.execute("DROP TRIGGER IF EXISTS trg_bibtext_update ON biblioteca_text;")
        c.execute("DROP TRIGGER IF EXISTS trg_bib_insert ON biblioteca;")
        c.execute("DROP TRIGGER IF EXISTS trg_bibtext_upsert ON biblioteca_text;")
        c.execute("DROP TRIGGER IF EXISTS trg_bib_upsert ON biblioteca;")
        c.execute("DROP FUNCTION IF EXISTS biblioteca_fts_update();")
        c.execute("DROP FUNCTION IF EXISTS biblioteca_fts_seed();")
        c.execute("DROP FUNCTION IF EXISTS biblioteca_fts_upsert();")

        # ---------- Função única com UPSERT idempotente ----------
        c.execute("""
        CREATE OR REPLACE FUNCTION biblioteca_fts_upsert() RETURNS trigger AS $$
        DECLARE
          v_conteudo  text := '';
          v_titulo    text := '';
          v_descricao text := '';
        BEGIN
          SELECT COALESCE(bt.conteudo,''), COALESCE(b.titulo,''), COALESCE(b.descricao,'')
            INTO v_conteudo, v_titulo, v_descricao
            FROM biblioteca b
            LEFT JOIN biblioteca_text bt ON bt.id = b.id
           WHERE b.id = NEW.id;

          INSERT INTO biblioteca_fts (id, doc)
               VALUES (NEW.id,
                       setweight(to_tsvector('portuguese', v_conteudo), 'A')
                    || setweight(to_tsvector('portuguese', v_titulo),   'B')
                    || setweight(to_tsvector('portuguese', v_descricao),'C'))
          ON CONFLICT (id) DO UPDATE
                SET doc = EXCLUDED.doc;

          RETURN NEW;
        END;
        $$ LANGUAGE plpgsql;
        """)

        # ---------- Triggers: ambos chamam a mesma função ----------
        c.execute("""
        CREATE TRIGGER trg_bibtext_upsert
        AFTER INSERT OR UPDATE ON biblioteca_text
        FOR EACH ROW
        EXECUTE FUNCTION biblioteca_fts_upsert();
        """)
        c.execute("""
        CREATE TRIGGER trg_bib_upsert
        AFTER INSERT OR UPDATE ON biblioteca
        FOR EACH ROW
        EXECUTE FUNCTION biblioteca_fts_upsert();
        """)

# ===================== Salvar conversa =====================
def salvar_conversa(usuario: str, pergunta: str, resposta: str):
    with conectar_banco() as conn, conn.cursor() as c:
        c.execute(
            "INSERT INTO conversas (usuario, pergunta, resposta) VALUES (%s, %s, %s)",
            (usuario, pergunta, resposta)
        )
        conn.commit()

# ===================== Extração de texto =====================
def extract_text_from_file(caminho: str) -> str:
    ext = os.path.splitext(caminho)[1].lower()
    try:
        if ext == '.pdf':
            from PyPDF2 import PdfReader
            reader = PdfReader(caminho)
            return '\n'.join(page.extract_text() or '' for page in reader.pages)
        if ext in ('.txt', '.md'):
            with open(caminho, encoding='utf-8') as f:
                return f.read()
        if ext == '.csv':
            import pandas as pd
            return pd.read_csv(caminho).to_csv(index=False)
        if ext in ('.xls', '.xlsx'):
            import pandas as pd
            return pd.read_excel(caminho).to_csv(index=False)
        if ext == '.docx':
            from docx import Document
            return '\n'.join(p.text for p in Document(caminho).paragraphs)
    except Exception:
        return ''
    return ''

# ===================== Helpers FTS =====================
def _fts_sanitize_query(q: str) -> str:
    return (q or "").strip()

def upsert_biblioteca_fts(id_: str, conteudo: str, titulo: str, descricao: str):
    with conectar_banco() as conn, conn.cursor() as c:
        c.execute("SELECT biblioteca_fts_update();")
        conn.commit()

# ===================== Registrar arquivo e indexar =====================
def registrar_arquivo(id_: str, titulo: str, descricao: str, tag: str, buffer):
    nome = f"{uuid4()}_{buffer.name}"
    caminho = os.path.join(UPLOAD_DIR, nome)
    with open(caminho, 'wb') as f:
        f.write(buffer.read())
    conteudo = extract_text_from_file(caminho)

    with conectar_banco() as conn, conn.cursor() as c:
        c.execute(
            "INSERT INTO biblioteca (id, titulo, descricao, tag, arquivo) VALUES (%s, %s, %s, %s, %s)",
            (id_, titulo, descricao, tag, nome)
        )
        c.execute(
            "INSERT INTO biblioteca_text (id, conteudo) VALUES (%s, %s)",
            (id_, conteudo)
        )
        conn.commit()

# ===================== Buscar procedimentos =====================
def buscar_procedimentos(
    query: Optional[str] = None,
    top_k: int = 5,
    truncate_chars: Optional[int] = None,
    mode: str = "full",
    fallback_chars: int = 2000,
) -> List[Tuple[str, str]]:
    with conectar_banco() as conn, conn.cursor() as c:
        rows: list[tuple] = []

        if query:
            q = _fts_sanitize_query(query)
            c.execute(
                """
                SELECT b.titulo, b.descricao, bt.conteudo
                  FROM biblioteca_fts f
                  JOIN biblioteca b ON b.id = f.id
                  JOIN biblioteca_text bt ON bt.id = b.id
                 WHERE f.doc @@ plainto_tsquery('portuguese', %s)
                 ORDER BY b.data_envio DESC
                 LIMIT %s
                """,
                (q, top_k)
            )
            rows = c.fetchall()

            if not rows:
                pattern = f"%{q}%"
                c.execute(
                    """
                    SELECT b.titulo, b.descricao, bt.conteudo
                      FROM biblioteca b
                      JOIN biblioteca_text bt ON bt.id = b.id
                     WHERE b.titulo ILIKE %s OR b.descricao ILIKE %s OR bt.conteudo ILIKE %s
                     ORDER BY b.data_envio DESC
                     LIMIT %s
                    """,
                    (pattern, pattern, pattern, top_k)
                )
                rows = c.fetchall()
        else:
            c.execute(
                """
                SELECT b.titulo, b.descricao, bt.conteudo
                  FROM biblioteca b
                  JOIN biblioteca_text bt ON bt.id = b.id
                 ORDER BY b.data_envio DESC
                 LIMIT %s
                """,
                (top_k,)
            )
            rows = c.fetchall()

    out: List[Tuple[str, str]] = []
    mode = (mode or "full").lower()
    for titulo, descricao, conteudo in rows:
        if mode == "desc":
            texto = (descricao or "").strip() or (conteudo or "")[:fallback_chars]
        elif mode == "mixed":
            desc_part = (descricao or "").strip()
            cont_part = (conteudo or "")[:fallback_chars]
            texto = desc_part + "\n\n" + cont_part if desc_part else cont_part
        else:  # full
            texto = conteudo or ""
            if truncate_chars and truncate_chars > 0:
                texto = texto[:truncate_chars]
        out.append((titulo, texto))
    return out

# ===================== Reindex manual (opcional) =====================
def reindex_biblioteca_from_uploads(default_tag: str = "Outro"):
    with conectar_banco() as conn, conn.cursor() as c:
        c.execute("DELETE FROM biblioteca_fts;")
        c.execute("DELETE FROM biblioteca_text;")
        c.execute("DELETE FROM biblioteca;")
        conn.commit()

        for fname in os.listdir(UPLOAD_DIR):
            fpath = os.path.join(UPLOAD_DIR, fname)
            if not os.path.isfile(fpath):
                continue
            doc_id = str(uuid4())
            titulo = os.path.splitext(fname)[0]
            descricao = ""
            tag = default_tag
            conteudo = extract_text_from_file(fpath)

            with conectar_banco() as conn2, conn2.cursor() as c2:
                c2.execute(
                    "INSERT INTO biblioteca (id, titulo, descricao, tag, arquivo) VALUES (%s, %s, %s, %s, %s)",
                    (doc_id, titulo, descricao, tag, fname)
                )
                c2.execute(
                    "INSERT INTO biblioteca_text (id, conteudo) VALUES (%s, %s)",
                    (doc_id, conteudo)
                )
                conn2.commit()

        with conectar_banco() as conn3, conn3.cursor() as c3:
            c3.execute("SELECT biblioteca_fts_update();")
            conn3.commit()


def buscar_procedimentos_prior_descricao(
    query: str,
    top_k_descricao: int = 3,
    top_k_full: int = 5,
    truncate_chars: int | None = None,
) -> list[tuple[str, str]]:
    if not query:
        return buscar_procedimentos(query=None, top_k=top_k_full, truncate_chars=truncate_chars)

    q = _fts_sanitize_query(query)
    with conectar_banco() as conn, conn.cursor() as c:
        c.execute(
            """
            SELECT b.id, b.titulo
              FROM biblioteca b
              JOIN biblioteca_fts f ON f.id = b.id
             WHERE setweight(to_tsvector('portuguese', COALESCE(b.titulo,'')), 'A')
                || setweight(to_tsvector('portuguese', COALESCE(b.descricao,'')), 'B')
                   @@ plainto_tsquery('portuguese', %s)
             ORDER BY b.data_envio DESC
             LIMIT %s
            """,
            (q, top_k_descricao)
        )
        cand_rows = c.fetchall()

        ids_escolhidos = [row[0] for row in cand_rows]

        if ids_escolhidos:
            ph = ",".join(["%s"] * len(ids_escolhidos))
            c.execute(
                f"""
                SELECT b.titulo, bt.conteudo
                  FROM biblioteca_text bt
                  JOIN biblioteca b ON b.id = bt.id
                 WHERE bt.id IN ({ph})
                 ORDER BY b.data_envio DESC
                """,
                ids_escolhidos
            )
            rows = c.fetchall()
        else:
            rows = buscar_procedimentos(query=q, top_k=top_k_full, truncate_chars=truncate_chars)

    if truncate_chars and truncate_chars > 0:
        rows = [(t, c[:truncate_chars]) for t, c in rows]

    return rows

